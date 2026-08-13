import openpyxl
import json
import os
from datetime import datetime

EXCEL_FILE = "VNIT Legacy Submission Form (Responses).xlsx"
JS_OUTPUT = "legacy_data.js"

if not os.path.exists(EXCEL_FILE):
    print(f"Error: Could not find '{EXCEL_FILE}'. Make sure it is in the same folder.")
    exit(1)

wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb.active

headers = [str(cell.value).strip() if cell.value else f"Col{i+1}" for i, cell in enumerate(ws[1])]

data = []
for row in ws.iter_rows(min_row=2, values_only=True):
    # Only include if they have a name and said Yes to being mentioned
    can_mention = str(row[4]).strip().lower() == 'yes' if row[4] else False
    if not can_mention:
        continue

    item = {}
    for i, h in enumerate(headers):
        val = row[i]
        if val is None:
            continue
        if isinstance(val, datetime):
            continue # skip timestamp
        val = str(val).strip()
        if not val or val.lower() == 'none':
            continue
        item[h] = val
    
    if "Name" in item:
        data.append(item)

# Write to JS file
js_content = f"const legacyData = {json.dumps(data, indent=4)};\n"
with open(JS_OUTPUT, "w", encoding="utf-8") as f:
    f.write(js_content)

print(f"Successfully processed {len(data)} responses and saved to {JS_OUTPUT}")
